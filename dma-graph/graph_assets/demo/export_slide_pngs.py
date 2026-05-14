"""
Export each PPTX slide as a PNG into a slides/ subfolder.
Reconstructs the slide by compositing the title bar over the embedded image.
"""

import os, io, textwrap
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SRC_PPTX   = os.path.join(SCRIPT_DIR, "slide4_clickthrough_full.pptx")
SLIDES_DIR = os.path.join(SCRIPT_DIR, "slides")

# Output resolution: 1920×1080
OUT_W, OUT_H = 1920, 1080
TITLE_H      = 42          # pixels

BG_COLOR     = (13,  17,  23)   # #0d1117
TITLE_BG     = (22,  27,  34)   # #161b22
ACC_COLOR    = (126, 58,  242)  # #7e3af2  — accent stripe + step_num
TEXT_COLOR   = (230, 237, 243)  # #e6edf3
MUT_COLOR    = (139, 148, 158)  # #8b949e


def get_font(size, bold=False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold
            else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf" if bold
            else "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def extract_picture(slide):
    """Return the PIL Image from the Picture shape on this slide."""
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img_bytes = shape.image.blob
            return Image.open(io.BytesIO(img_bytes)).convert("RGB")
    return None


def render_slide_png(slide, out_w=OUT_W, out_h=OUT_H, title_h=TITLE_H):
    """Composite title bar + embedded image into a single PNG."""
    canvas = Image.new("RGB", (out_w, out_h), BG_COLOR)
    draw   = ImageDraw.Draw(canvas)

    # ── Title bar background ──────────────────────────────────────────────────
    draw.rectangle([(0, 0), (out_w, title_h)], fill=TITLE_BG)

    # ── Left accent stripe ────────────────────────────────────────────────────
    draw.rectangle([(0, 0), (5, title_h)], fill=ACC_COLOR)

    # ── Step label + title text ───────────────────────────────────────────────
    # Collect text runs from the single textbox
    step_text  = ""
    title_text = ""
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                runs = para.runs
                if len(runs) >= 2:
                    step_text  = runs[0].text.strip()
                    title_text = runs[1].text.strip()
                elif len(runs) == 1:
                    title_text = runs[0].text.strip()
            break   # only one textbox per slide

    font_step  = get_font(14, bold=True)
    font_title = get_font(14, bold=True)

    x = 14
    y = (title_h - 16) // 2

    # Draw step_num in accent colour
    if step_text:
        draw.text((x, y), step_text + "  ", font=font_step, fill=ACC_COLOR)
        bbox = draw.textbbox((x, y), step_text + "  ", font=font_step)
        x = bbox[2]

    # Draw title in light colour
    draw.text((x, y), title_text, font=font_title, fill=TEXT_COLOR)

    # ── Embedded image fills remaining height ─────────────────────────────────
    img = extract_picture(slide)
    if img:
        img = img.resize((out_w, out_h - title_h), Image.LANCZOS)
        canvas.paste(img, (0, title_h))

    return canvas


def main():
    os.makedirs(SLIDES_DIR, exist_ok=True)
    prs = Presentation(SRC_PPTX)

    for i, slide in enumerate(prs.slides, start=1):
        png = render_slide_png(slide)
        out_path = os.path.join(SLIDES_DIR, f"slide_{i:02d}.png")
        png.save(out_path, format="PNG", optimize=True)
        print(f"  {out_path}  ({os.path.getsize(out_path)//1024} KB)")

    print(f"\nDone — {len(prs.slides)} PNGs in {SLIDES_DIR}/")


if __name__ == "__main__":
    main()
